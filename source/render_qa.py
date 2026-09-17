"""Render each slide of a .pptx to PNG using python-pptx + PIL.

LibreOffice cannot load documents in this sandbox, so this stands in for
`soffice --convert-to pdf` during visual QA. It is an approximation: text is
wrapped with real Noto Sans CJK metrics, but effects and gradients are ignored.
"""
import sys, os, io
from pptx import Presentation
from pptx.util import Emu
from PIL import Image, ImageDraw, ImageFont

DPI = 150
EMU_IN = 914400
REG = "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"
BLD = "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc"
NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}
_fc = {}


def font(sz, bold):
    k = (round(sz, 1), bold)
    if k not in _fc:
        _fc[k] = ImageFont.truetype(BLD if bold else REG, max(1, int(round(sz * DPI / 72))))
    return _fc[k]


def px(emu):
    return emu / EMU_IN * DPI


def hexrgb(s, default=(0, 0, 0)):
    try:
        return tuple(int(s[i:i + 2], 16) for i in (0, 2, 4))
    except Exception:
        return default


def blend(c, bg, alpha):
    return tuple(int(c[i] * alpha + bg[i] * (1 - alpha)) for i in range(3))


def solid_of(el, deep=False):
    """(rgb, alpha) from el's solidFill, or None. Direct child unless deep."""
    sf = el.find(".//a:solidFill", NS) if deep else el.find("a:solidFill", NS)
    if sf is None:
        return None
    clr = sf.find("a:srgbClr", NS)
    if clr is None:
        return None
    rgb = hexrgb(clr.get("val"))
    a = clr.find("a:alpha", NS)
    alpha = int(a.get("val")) / 100000 if a is not None else 1.0
    return rgb, alpha


def wrap(text, f, maxw):
    out = []
    for para in text.split("\n"):
        if para == "":
            out.append("")
            continue
        cur, i = "", 0
        while i < len(para):
            ch = para[i]
            t = cur + ch
            if f.getlength(t) <= maxw or cur == "":
                cur, i = t, i + 1
            elif ord(ch) < 0x1100 and " " in cur.strip():
                bp = cur.rstrip().rfind(" ")
                out.append(cur[:bp])
                cur = cur[bp + 1:]
            else:
                out.append(cur)
                cur = ""
        if cur:
            out.append(cur)
    return out


def draw_background(slide, prs, img, dr, Wp, Hp):
    bg = slide._element.find("p:cSld/p:bg", NS)
    if bg is None:
        dr.rectangle([0, 0, Wp, Hp], fill=(255, 255, 255))
        return (255, 255, 255)
    blip = bg.find(".//a:blip", NS)
    if blip is not None:
        rid = blip.get("{%s}embed" % NS["r"])
        part = slide.part.related_part(rid)
        bim = Image.open(io.BytesIO(part.blob)).convert("RGB").resize((Wp, Hp), Image.LANCZOS)
        img.paste(bim, (0, 0))
        return (40, 40, 40)
    s = solid_of(bg, deep=True)
    col = s[0] if s else (255, 255, 255)
    dr.rectangle([0, 0, Wp, Hp], fill=col)
    return col


def render(path, outdir):
    prs = Presentation(path)
    Wp, Hp = int(px(prs.slide_width)), int(px(prs.slide_height))
    os.makedirs(outdir, exist_ok=True)
    issues = []

    for idx, slide in enumerate(prs.slides, 1):
        img = Image.new("RGB", (Wp, Hp), "white")
        dr = ImageDraw.Draw(img)
        bgcol = draw_background(slide, prs, img, dr, Wp, Hp)

        for sh in slide.shapes:
            if sh.left is None:
                continue
            L, T = px(sh.left), px(sh.top)
            Wd, Ht = px(sh.width), px(sh.height)
            if L < -1 or T < -1 or L + Wd > Wp + 1 or T + Ht > Hp + 1:
                issues.append(f"[{idx}] off-slide: {sh.shape_type} "
                              f"x={L/DPI:.2f} y={T/DPI:.2f} w={Wd/DPI:.2f} h={Ht/DPI:.2f}")

            if sh.shape_type == 13:  # PICTURE
                try:
                    pim = Image.open(io.BytesIO(sh.image.blob)).convert("RGB")
                    img.paste(pim.resize((max(1, int(Wd)), max(1, int(Ht))), Image.LANCZOS),
                              (int(L), int(T)))
                except Exception as e:
                    issues.append(f"[{idx}] image fail: {e}")
                continue

            el = sh._element
            spPr = el.find(".//p:spPr", NS)
            geom = "rect"
            if spPr is not None:
                pg = spPr.find("a:prstGeom", NS)
                if pg is not None:
                    geom = pg.get("prst")

            fill = solid_of(spPr) if spPr is not None else None
            ln_el = spPr.find("a:ln", NS) if spPr is not None else None
            lncol = None
            if ln_el is not None and ln_el.find("a:noFill", NS) is None:
                s = solid_of(ln_el)
                if s:
                    lncol = s[0]
            box = [int(L), int(T), int(L + Wd), int(T + Ht)]

            # composite through an RGBA layer so translucent fills blend with what is
            # actually underneath (background photo, overlapping circles) rather than
            # with a single assumed colour
            if geom == "line":
                if lncol:
                    dr.line([int(L), int(T), int(L + Wd), int(T + Ht)], fill=lncol, width=2)
            else:
                layer = Image.new("RGBA", img.size, (0, 0, 0, 0))
                ld = ImageDraw.Draw(layer)
                shape_fn = ld.ellipse if geom == "ellipse" else ld.rectangle
                if fill:
                    shape_fn(box, fill=fill[0] + (int(fill[1] * 255),))
                if lncol:
                    shape_fn(box, outline=lncol + (255,), width=2)
                img.paste(Image.alpha_composite(img.convert("RGBA"), layer).convert("RGB"), (0, 0))
                dr = ImageDraw.Draw(img)

            if not sh.has_text_frame or not sh.text_frame.text.strip():
                continue

            tf = sh.text_frame
            anchor = (tf.vertical_anchor.__str__() if tf.vertical_anchor is not None else "TOP")
            lines = []
            for para in tf.paragraphs:
                runs = [r for r in para.runs if r.text]
                if not runs:
                    lines.append(("", 11, False, (0, 0, 0), None))
                    continue
                r0 = runs[0]
                sz = r0.font.size.pt if r0.font.size else 11
                bold = bool(r0.font.bold)
                col = (0, 0, 0)
                try:
                    if r0.font.color is not None and r0.font.color.rgb is not None:
                        col = hexrgb(str(r0.font.color.rgb))
                except Exception:
                    pass
                ls = para.line_spacing
                lsp = ls.pt if hasattr(ls, "pt") else (sz * 1.2 if ls is None else sz * ls)
                f = font(sz, bold)
                al = str(para.alignment) if para.alignment is not None else "LEFT"
                txt = "".join(r.text for r in runs)
                for ln in wrap(txt, f, Wd):
                    lines.append((ln, sz, bold, col, (lsp, al)))

            total = sum((l[4][0] if l[4] else l[1] * 1.2) * DPI / 72 for l in lines)
            if total > Ht + 4 and total - Ht > 0.10 * DPI:
                issues.append(f"[{idx}] text overflow {total/DPI:.2f}in > box {Ht/DPI:.2f}in "
                              f"@ x={L/DPI:.2f} y={T/DPI:.2f} :: {sh.text_frame.text[:44]!r}")

            y = T
            if "MIDDLE" in anchor:
                y = T + (Ht - total) / 2
            elif "BOTTOM" in anchor:
                y = T + Ht - total
            for ln, sz, bold, col, meta in lines:
                lsp, al = meta if meta else (sz * 1.2, "LEFT")
                f = font(sz, bold)
                w = f.getlength(ln)
                x = L
                if "CENTER" in al:
                    x = L + (Wd - w) / 2
                elif "RIGHT" in al:
                    x = L + Wd - w
                asc = f.getmetrics()[0]
                dr.text((x, y + (lsp * DPI / 72 - asc) / 2), ln, font=f, fill=col)
                y += lsp * DPI / 72

        img.save(os.path.join(outdir, f"slide-{idx:02d}.png"))

    print(f"rendered {len(prs.slides)} slides -> {outdir}")
    if issues:
        print(f"\n{len(issues)} issue(s):")
        for i in issues:
            print("  " + i)
    else:
        print("no bounds/overflow issues detected")


if __name__ == "__main__":
    render(sys.argv[1], sys.argv[2])
